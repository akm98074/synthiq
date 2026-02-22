"""
DOCX parser using python-docx.

Extracts paragraph text and groups it into synthetic "pages" so the pipeline
treats DOCX like any other multi-page document.
"""
from __future__ import annotations

import io

from docx import Document  # type: ignore[import-untyped]

from .base import PageText

_PARAGRAPHS_PER_PAGE = 50  # Approximate grouping for DOCX "pages"


def parse_docx(content: bytes) -> list[PageText]:
    """
    Parse a DOCX file from raw bytes.

    Returns PageText objects where each "page" contains up to 50 paragraphs.
    """
    doc = Document(io.BytesIO(content))

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Also include text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    paragraphs.append(cell_text)

    if not paragraphs:
        return []

    pages: list[PageText] = []
    for i in range(0, len(paragraphs), _PARAGRAPHS_PER_PAGE):
        group = paragraphs[i : i + _PARAGRAPHS_PER_PAGE]
        pages.append(
            PageText(
                page_number=i // _PARAGRAPHS_PER_PAGE + 1,
                text="\n\n".join(group),
            )
        )

    return pages
